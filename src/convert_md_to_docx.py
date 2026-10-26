"""
src/convert_md_to_docx.py
Converts docs/RESEARCH_PAPER_FULL_MANUSCRIPT.md into a beautifully formatted Word (.docx) document
with standard Academic Journal styling (1-inch margins, Times New Roman, formatted tables,
embedded high-resolution 300 DPI figures, centered captions, callout panels).
"""

import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MD_PATH = os.path.join(PROJECT_ROOT, "docs", "RESEARCH_PAPER_FULL_MANUSCRIPT.md")
DOCX_PATH = os.path.join(PROJECT_ROOT, "docs", "RESEARCH_PAPER_FULL_MANUSCRIPT.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def add_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CCCCCC"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_callout_box(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F8FAFC")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=160)
    
    # Left border only (Callout style)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="1E3A8A"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    doc.add_paragraph() # Spacing

def resolve_image_path(img_ref):
    # Relative to docs/ directory
    docs_dir = os.path.join(PROJECT_ROOT, "docs")
    clean_ref = img_ref.strip()
    candidate = os.path.normpath(os.path.join(docs_dir, clean_ref))
    if os.path.exists(candidate):
        return candidate
    candidate2 = os.path.normpath(os.path.join(PROJECT_ROOT, clean_ref))
    if os.path.exists(candidate2):
        return candidate2
    return None

def build_docx():
    print(f"Reading Markdown from: {MD_PATH}")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = docx.Document()
    
    # Page setup: Standard Letter, 1-inch margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(4)

    in_table = False
    table_lines = []

    def flush_table(lines_to_flush):
        if not lines_to_flush:
            return
        rows_data = []
        for line in lines_to_flush:
            line_str = line.strip()
            if not line_str.startswith("|"):
                continue
            cols = [c.strip() for c in line_str.split("|")[1:-1]]
            # Skip separator line (e.g. |---|---|)
            if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                continue
            rows_data.append(cols)

        if not rows_data:
            return

        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=len(rows_data), cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        add_table_borders(table)

        for i, row in enumerate(rows_data):
            for j in range(num_cols):
                val = row[j] if j < len(row) else ""
                cell = table.cell(i, j)
                set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(9.5)
                if i == 0:
                    set_cell_background(cell, "1E3A8A") # Navy header
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF) # White text
                else:
                    if i % 2 == 1:
                        set_cell_background(cell, "F8FAFC") # Soft zebra
                    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

        doc.add_paragraph()  # Spacing after table

    i = 0
    images_embedded = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")

        # Table detection
        if line.strip().startswith("|") and "|" in line.strip()[1:]:
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            flush_table(table_lines)
            table_lines = []
            in_table = False

        stripped = line.strip()

        # Blank line
        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped in ["---", "***", "___"]:
            i += 1
            continue

        # Image detection: ![Caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_rel = img_match.group(2)
            img_abs = resolve_image_path(img_rel)
            if img_abs and os.path.exists(img_abs):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(12)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(img_abs, width=Inches(5.8))
                images_embedded += 1
                print(f"  [+] Embedded image: {img_rel} -> {os.path.basename(img_abs)}")
            else:
                print(f"  [!] Image not found: {img_rel}")
            i += 1
            continue

        # Figure Caption detection: *Figure X: ...*
        if stripped.startswith("*Figure ") and stripped.endswith("*"):
            cap_text = stripped[1:-1].strip()
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_before = Pt(2)
            p_cap.paragraph_format.space_after = Pt(12)
            run_cap = p_cap.add_run(cap_text)
            run_cap.italic = True
            run_cap.font.name = 'Times New Roman'
            run_cap.font.size = Pt(9.5)
            run_cap.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
            i += 1
            continue

        # Blockquote / Callout box: > ...
        if stripped.startswith("> "):
            callout_text = stripped[2:].strip()
            # Gather consecutive quote lines
            quote_lines = [callout_text]
            while i + 1 < len(lines) and lines[i+1].strip().startswith(">"):
                i += 1
                q_line = lines[i].strip()
                if q_line.startswith("> "):
                    quote_lines.append(q_line[2:].strip())
                elif q_line == ">":
                    quote_lines.append("\n")
            full_quote = " ".join(quote_lines).replace(" \n ", "\n\n")
            add_callout_box(doc, full_quote)
            i += 1
            continue

        # Title: # ...
        if stripped.startswith("# "):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(title_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(17)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Deep Navy
            i += 1
            continue

        # Heading 1: ## ...
        if stripped.startswith("## "):
            h2_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h2_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13.5)
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Academic Navy
            i += 1
            continue

        # Heading 2: ### ...
        if stripped.startswith("### "):
            h3_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(h3_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11.5)
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
            i += 1
            continue

        # Heading 3: #### ...
        if stripped.startswith("#### "):
            h4_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(h4_text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            i += 1
            continue

        # Bullet list items
        if stripped.startswith("- ") or stripped.startswith("* "):
            item_text = stripped[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Numbered list items
        num_match = re.match(r'^(\d+)\.\s+(.*)$', stripped)
        if num_match:
            item_text = num_match.group(2)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', item_text)
            for part in parts:
                if part.startswith("**") and part.endswith("**"):
                    r = p.add_run(part[2:-2])
                    r.bold = True
                else:
                    p.add_run(part)
            i += 1
            continue

        # Code block / ASCII architecture diagrams
        if stripped.startswith("```"):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            code_text = "\n".join(code_lines)
            p_code = doc.add_paragraph()
            p_code.paragraph_format.space_before = Pt(6)
            p_code.paragraph_format.space_after = Pt(6)
            run_code = p_code.add_run(code_text)
            run_code.font.name = 'Consolas'
            run_code.font.size = Pt(8.5)
            run_code.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*|\*.*?\*)', stripped)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                r = p.add_run(part[2:-2])
                r.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 1:
                r = p.add_run(part[1:-1])
                r.italic = True
            else:
                p.add_run(part)

        i += 1

    if in_table:
        flush_table(table_lines)

    doc.save(DOCX_PATH)
    file_size_mb = os.path.getsize(DOCX_PATH) / (1024 * 1024)
    print(f"Successfully generated DOCX with {images_embedded} images ({file_size_mb:.2f} MB) at: {DOCX_PATH}")

if __name__ == "__main__":
    build_docx()
